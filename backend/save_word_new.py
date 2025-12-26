def save_word(file_path: Path, content: Dict[str, Any], template_filename: str = None):
    """
    Word dosyasını kaydet - orijinal formatı koruyarak
    Şablon dosyayı kopyalar ve sadece metin içeriğini günceller
    """
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    # Şablon dosya yolunu belirle
    if template_filename:
        template_path = TEMPLATE_DIR / template_filename
    else:
        template_path = None
    
    # Şablon varsa kopyala, yoksa yeni belge oluştur
    if template_path and template_path.exists():
        # Şablon dosyayı hedefe kopyala (formatı korumak için)
        shutil.copy(template_path, file_path)
        
        # Kopyalanan dosyayı aç
        doc = Document(file_path)
        
        # Şimdi sadece metin içeriğini güncelle
        paragraphs_data = content.get("paragraphs", [])
        tables_data = content.get("tables", [])
        
        # Paragrafları güncelle (sadece metin içeriği)
        doc_paragraphs = [p for p in doc.paragraphs if p.text.strip()]  # Boş olmayan paragraflar
        for i, para_data in enumerate(paragraphs_data):
            if i < len(doc_paragraphs):
                # Mevcut paragrafın metnini güncelle, formatları koru
                p = doc_paragraphs[i]
                new_text = para_data.get("text", "")
                
                # Tüm runs'ları temizle ve yeni metni ekle
                # İlk run'ı kullan ve formatını koru
                if p.runs:
                    # İlk run'ın formatını kullan
                    first_run = p.runs[0]
                    # Tüm runs'ları temizle
                    for run in p.runs:
                        run.text = ""
                    # İlk run'a yeni metni yaz
                    first_run.text = new_text
                else:
                    # Run yoksa metni direkt ekle
                    p.text = new_text
        
        # Tabloları güncelle (sadece metin içeriği)
        for table_idx, table_data in enumerate(tables_data):
            if table_idx < len(doc.tables):
                table = doc.tables[table_idx]
                
                for row_idx, row_data in enumerate(table_data):
                    if row_idx < len(table.rows):
                        row = table.rows[row_idx]
                        
                        for col_idx, cell_data in enumerate(row_data):
                            if col_idx < len(row.cells):
                                cell = row.cells[col_idx]
                                
                                # Cell içeriğini güncelle
                                if isinstance(cell_data, dict):
                                    new_text = cell_data.get("text", "")
                                else:
                                    new_text = str(cell_data)
                                
                                # Cell'in ilk paragrafını güncelle
                                if cell.paragraphs:
                                    cell_para = cell.paragraphs[0]
                                    if cell_para.runs:
                                        # İlk run'ın formatını koru
                                        first_run = cell_para.runs[0]
                                        for run in cell_para.runs:
                                            run.text = ""
                                        first_run.text = new_text
                                    else:
                                        cell_para.text = new_text
                                else:
                                    cell.text = new_text
        
        doc.save(file_path)
        
    else:
        # Şablon yoksa yeni belge oluştur (eski yöntem)
        doc = Document()
        
        # Paragrafları ekle
        paragraphs = content.get("paragraphs", [])
        for para_data in paragraphs:
            para = doc.add_paragraph()
            
            # Paragraph stilini ayarla
            style_name = para_data.get("style", "Normal")
            try:
                para.style = style_name
            except:
                para.style = "Normal"
            
            # Metni ekle
            text = para_data.get("text", "")
            
            # Format bilgileri varsa uygula
            if "format" in para_data:
                format_data = para_data["format"]
                run = para.add_run(text)
                
                # Font stilleri
                if format_data.get("bold"):
                    run.bold = True
                if format_data.get("italic"):
                    run.italic = True
                if format_data.get("underline"):
                    run.underline = True
                
                # Font boyutu
                if format_data.get("font_size"):
                    run.font.size = Pt(format_data["font_size"])
                
                # Font adı
                if format_data.get("font_name"):
                    run.font.name = format_data["font_name"]
                
                # Font rengi
                if format_data.get("font_color"):
                    try:
                        # RGB hex string'den RGBColor'a çevir
                        rgb_str = format_data["font_color"]
                        if len(rgb_str) == 8:  # ARGB formatı
                            rgb_str = rgb_str[2:]  # Alpha'yı atla
                        if len(rgb_str) == 6:
                            r = int(rgb_str[0:2], 16)
                            g = int(rgb_str[2:4], 16)
                            b = int(rgb_str[4:6], 16)
                            run.font.color.rgb = RGBColor(r, g, b)
                    except:
                        pass
            else:
                # Format bilgisi yoksa sadece metni ekle
                para.add_run(text)
            
            # Hizalama
            if "alignment" in para_data:
                align_str = para_data["alignment"]
                if "LEFT" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif "CENTER" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif "RIGHT" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif "JUSTIFY" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Tabloları ekle
        tables = content.get("tables", [])
        for table_data in tables:
            if not table_data or not table_data[0]:
                continue
            
            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            for i, row_data in enumerate(table_data):
                for j, cell_data in enumerate(row_data):
                    # Cell_data bir dict veya string olabilir
                    if isinstance(cell_data, dict):
                        cell_text = cell_data.get("text", "")
                        cell = table.cell(i, j)
                        cell.text = cell_text
                        
                        # Format bilgileri varsa uygula
                        if "format" in cell_data:
                            format_data = cell_data["format"]
                            if cell.paragraphs:
                                para = cell.paragraphs[0]
                                run = para.runs[0] if para.runs else para.add_run(cell_text)
                                
                                # Font stilleri
                                if format_data.get("bold"):
                                    run.bold = True
                                if format_data.get("italic"):
                                    run.italic = True
                                
                                # Font boyutu
                                if format_data.get("font_size"):
                                    run.font.size = Pt(format_data["font_size"])
                                
                                # Font adı
                                if format_data.get("font_name"):
                                    run.font.name = format_data["font_name"]
                                
                                # Font rengi
                                if format_data.get("font_color"):
                                    try:
                                        rgb_str = format_data["font_color"]
                                        if len(rgb_str) == 8:
                                            rgb_str = rgb_str[2:]
                                        if len(rgb_str) == 6:
                                            r = int(rgb_str[0:2], 16)
                                            g = int(rgb_str[2:4], 16)
                                            b = int(rgb_str[4:6], 16)
                                            run.font.color.rgb = RGBColor(r, g, b)
                                    except:
                                        pass
                    else:
                        # Eski format (sadece string)
                        table.cell(i, j).text = str(cell_data)
        
        doc.save(file_path)
