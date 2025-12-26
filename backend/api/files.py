from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from docx import Document
import json
import shutil
import os

router = APIRouter()

TEMPLATE_DIR = Path("form_sablonlari")
COMPANIES_DIR = Path("firmalar")
COMPANIES_DIR.mkdir(exist_ok=True)

@router.get("/file/{filename}/raw")
async def get_file_raw(filename: str):
    """
    Şablon dosyasını binary olarak döndür (Word görüntüleme için)
    """
    file_path = TEMPLATE_DIR / filename
    
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    
    return FileResponse(
        path=file_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename
    )

@router.get("/file/{filename}")
async def get_file(filename: str):
    """
    Şablon dosyasını oku
    """
    file_path = TEMPLATE_DIR / filename
    
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    
    try:
        extension = file_path.suffix.lower()
        
        if extension in ['.xlsx', '.xls']:
            return read_excel(file_path)
        elif extension in ['.docx', '.doc']:
            return read_word(file_path)
        else:
            raise HTTPException(status_code=400, detail="Desteklenmeyen dosya formatı")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Dosya okuma hatası: {str(e)}")

@router.get("/companies/{company_name}/file/{filename}/raw")
async def get_company_file_raw(company_name: str, filename: str):
    """
    Firma dosyasını binary olarak döndür (Word görüntüleme için)
    """
    company_dir = COMPANIES_DIR / company_name
    
    if not company_dir.exists():
        raise HTTPException(status_code=404, detail="Firma bulunamadı")
    
    # Dosyayı firma klasöründe veya alt klasörlerde ara
    file_path = None
    for found_path in company_dir.rglob(filename):
        if found_path.is_file() and found_path.name == filename:
            file_path = found_path
            break
    
    if not file_path or not file_path.exists():
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    
    return FileResponse(
        path=file_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename
    )

@router.get("/companies/{company_name}/file/{filename}")
async def get_company_file(company_name: str, filename: str):
    """
    Firma dosyasını oku (alt klasörler dahil)
    """
    company_dir = COMPANIES_DIR / company_name
    
    if not company_dir.exists():
        raise HTTPException(status_code=404, detail="Firma bulunamadı")
    
    # Dosyayı firma klasöründe veya alt klasörlerde ara
    file_path = None
    for found_path in company_dir.rglob(filename):
        if found_path.is_file() and found_path.name == filename:
            file_path = found_path
            break
    
    if not file_path or not file_path.exists():
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    
    try:
        extension = file_path.suffix.lower()
        
        if extension in ['.xlsx', '.xls']:
            return read_excel(file_path)
        elif extension in ['.docx', '.doc']:
            return read_word(file_path)
        else:
            raise HTTPException(status_code=400, detail="Desteklenmeyen dosya formatı")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Dosya okuma hatası: {str(e)}")

@router.delete("/file/{filename}")
async def delete_file(filename: str):
    """
    Şablon dosyasını sil
    """
    file_path = TEMPLATE_DIR / filename
    
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    
    try:
        os.remove(file_path)
        return {
            "message": "Dosya başarıyla silindi",
            "filename": filename
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Dosya silme hatası: {str(e)}")

@router.delete("/companies/{company_name}/file/{filename}")
async def delete_company_file(company_name: str, filename: str):
    """
    Firma dosyasını sil (alt klasörler dahil)
    """
    company_dir = COMPANIES_DIR / company_name
    
    if not company_dir.exists():
        raise HTTPException(status_code=404, detail="Firma bulunamadı")
    
    # Dosyayı firma klasöründe veya alt klasörlerde ara
    file_path = None
    for found_path in company_dir.rglob(filename):
        if found_path.is_file() and found_path.name == filename:
            file_path = found_path
            break
    
    if not file_path or not file_path.exists():
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    
    try:
        os.remove(file_path)
        return {
            "message": "Dosya başarıyla silindi",
            "company": company_name,
            "filename": filename
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Dosya silme hatası: {str(e)}")

def read_excel(file_path: Path) -> Dict[str, Any]:
    """Excel dosyasını oku ve tüm format bilgileriyle JSON formatına çevir"""
    wb = openpyxl.load_workbook(file_path)
    
    sheets_data = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Sheet'teki tüm verileri ve stilleri al
        data = []
        for row in ws.iter_rows():
            row_data = []
            for cell in row:
                # Hücre değeri
                value = cell.value if cell.value is not None else ""
                
                # Hücre stilleri
                cell_style = {
                    "value": str(value),
                    "coordinate": cell.coordinate,
                }
                
                # Font stilleri
                if cell.font:
                    cell_style["font"] = {
                        "bold": cell.font.bold or False,
                        "italic": cell.font.italic or False,
                        "size": cell.font.size or 11,
                        "color": cell.font.color.rgb if cell.font.color and hasattr(cell.font.color, 'rgb') else None,
                        "name": cell.font.name or "Calibri"
                    }
                
                # Arka plan rengi
                if cell.fill and cell.fill.start_color:
                    if hasattr(cell.fill.start_color, 'rgb'):
                        cell_style["fill"] = cell.fill.start_color.rgb
                
                # Hizalama
                if cell.alignment:
                    cell_style["alignment"] = {
                        "horizontal": cell.alignment.horizontal or "left",
                        "vertical": cell.alignment.vertical or "top",
                        "wrap_text": cell.alignment.wrap_text or False
                    }
                
                # Kenarlıklar
                if cell.border:
                    cell_style["border"] = {
                        "top": bool(cell.border.top and cell.border.top.style),
                        "bottom": bool(cell.border.bottom and cell.border.bottom.style),
                        "left": bool(cell.border.left and cell.border.left.style),
                        "right": bool(cell.border.right and cell.border.right.style)
                    }
                
                row_data.append(cell_style)
            data.append(row_data)
        
        # Birleştirilmiş hücreler
        merged_cells = []
        for merged_range in ws.merged_cells.ranges:
            merged_cells.append({
                "start_row": merged_range.min_row,
                "start_col": merged_range.min_col,
                "end_row": merged_range.max_row,
                "end_col": merged_range.max_col,
                "range": str(merged_range)
            })
        
        # Sütun genişlikleri
        column_widths = {}
        for col_letter, col_dimension in ws.column_dimensions.items():
            if col_dimension.width:
                column_widths[col_letter] = col_dimension.width
        
        # Satır yükseklikleri
        row_heights = {}
        for row_num, row_dimension in ws.row_dimensions.items():
            if row_dimension.height:
                row_heights[str(row_num)] = row_dimension.height
        
        sheets_data[sheet_name] = {
            "data": data,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_cells": merged_cells,
            "column_widths": column_widths,
            "row_heights": row_heights
        }
    
    return {
        "type": "excel",
        "filename": file_path.name,
        "sheets": sheets_data,
        "active_sheet": wb.active.title
    }

def read_word(file_path: Path) -> Dict[str, Any]:
    """Word dosyasını oku ve JSON formatına çevir - stil bilgileriyle birlikte"""
    doc = Document(file_path)
    
    # Paragrafları oku
    paragraphs = []
    for para in doc.paragraphs:
        para_data = {
            "text": para.text,
            "style": para.style.name if para.style else "Normal"
        }
        
        # Paragraph format bilgileri
        if para.runs:
            # İlk run'un formatını al (genelde tüm paragraph aynı formatta)
            first_run = para.runs[0]
            para_data["format"] = {
                "bold": first_run.bold if first_run.bold is not None else False,
                "italic": first_run.italic if first_run.italic is not None else False,
                "underline": bool(first_run.underline),
                "font_size": first_run.font.size.pt if first_run.font.size else None,
                "font_name": first_run.font.name if first_run.font.name else None,
                "font_color": str(first_run.font.color.rgb) if first_run.font.color and hasattr(first_run.font.color, 'rgb') else None
            }
            
            # Alignment
            if para.alignment:
                para_data["alignment"] = str(para.alignment)
        
        paragraphs.append(para_data)
    
    # Tabloları oku
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                cell_data = {
                    "text": cell.text
                }
                # Cell format bilgileri
                if cell.paragraphs:
                    first_para = cell.paragraphs[0]
                    if first_para.runs:
                        first_run = first_para.runs[0]
                        cell_data["format"] = {
                            "bold": first_run.bold if first_run.bold is not None else False,
                            "italic": first_run.italic if first_run.italic is not None else False,
                            "font_size": first_run.font.size.pt if first_run.font.size else None,
                            "font_name": first_run.font.name if first_run.font.name else None,
                            "font_color": str(first_run.font.color.rgb) if first_run.font.color and hasattr(first_run.font.color, 'rgb') else None
                        }
                row_data.append(cell_data)
            table_data.append(row_data)
        tables.append(table_data)
    
    return {
        "type": "word",
        "filename": file_path.name,
        "paragraphs": paragraphs,
        "tables": tables
    }

def sanitize_folder_name(name: str) -> str:
    """
    Klasör adını temizle - geçersiz karakterleri kaldır
    """
    import re
    name = re.sub(r'[<>:"/\\|?*]', '_', str(name))
    name = name.strip()
    return name if name else None

def extract_work_order_number(sheets_data: Dict[str, Any]) -> str:
    """
    Excel içeriğinden iş emri numarasını çıkar
    SM ile başlayan herhangi bir hücreyi iş emri no olarak kabul eder
    Fixed: Removed emoji characters for Windows compatibility
    """
    import re
    
    print(f"\n====== EXTRACT_WORK_ORDER_NUMBER (SM ile başlayan hücre aranıyor) ======")
    print(f"Sheets: {list(sheets_data.keys())}")
    
    try:
        for sheet_name, sheet in sheets_data.items():
            data = sheet.get('data', [])
            
            print(f"\n--- Sheet: {sheet_name} ---")
            
            for row_idx, row in enumerate(data):
                for col_idx, cell in enumerate(row):
                    cell_value = str(cell.get('value', '')).strip()
                    
                    if not cell_value:
                        continue
                    
                    # SM ile başlıyor mu?
                    if cell_value.upper().startswith('SM'):
                        print(f"[OK] Is Emri NO bulundu: Row {row_idx}, Col {col_idx} = '{cell_value}'")
                        sanitized = sanitize_folder_name(cell_value)
                        if sanitized:
                            print(f"   [INFO] IS EMRI NO: '{sanitized}'")
                            return sanitized
        
        print("[WARN] SM ile baslayan hucre bulunamadi")
        return None
    except Exception as e:
        print(f"[ERROR] Hata: {str(e)}")
        import traceback
        traceback.print_exc()
        return None

@router.post("/save")
async def save_file(data: Dict[str, Any]):
    """
    Düzenlenmiş dosyayı firma klasörüne kaydet
    Dosya adı başına firma ismi eklenir (örn: BAYKAR_F.02.xlsx)
    İş emri no bulunursa alt klasör oluşturulur (örn: firmalar/Baykar/SM-128/)
    """
    try:
        filename = data.get("filename")
        company_name = data.get("company")
        file_type = data.get("type")
        content = data.get("content")
        
        if not all([filename, company_name, file_type, content]):
            raise HTTPException(status_code=400, detail="Eksik parametre")
        
        # Firma klasörünü oluştur
        company_dir = COMPANIES_DIR / company_name
        company_dir.mkdir(exist_ok=True)
        
        # Excel ise iş emri numarasını çıkar
        work_order_no = None
        if file_type == "excel" and "sheets" in content:
            work_order_no = extract_work_order_number(content["sheets"])
        
        # Hedef klasör yolu
        if work_order_no:
            # İş emri no varsa alt klasör oluştur
            target_dir = company_dir / work_order_no
            target_dir.mkdir(parents=True, exist_ok=True)
            print(f"Klasör oluşturuldu: {target_dir}")
        else:
            # İş emri no yoksa doğrudan firma klasörüne kaydet
            target_dir = company_dir
            print(f"Doğrudan firma klasörüne kaydediliyor")
        
        # Dosya adına firma ismini ekle (büyük harflerle)
        company_prefix = company_name.upper().replace(' ', '_')
        new_filename = f"{company_prefix}_{filename}"
        
        # Hedef dosya yolu
        target_path = target_dir / new_filename
        
        # Dosya tipine göre kaydet
        if file_type == "excel":
            save_excel(target_path, content, filename)
        elif file_type == "word":
            save_word(target_path, content, filename)  # filename'i template olarak gönder
        else:
            raise HTTPException(status_code=400, detail="Desteklenmeyen dosya tipi")
        
        return {
            "message": "Dosya başarıyla kaydedildi",
            "path": str(target_path),
            "company": company_name,
            "work_order_no": work_order_no,
            "filename": new_filename
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Dosya kaydetme hatası: {str(e)}")

def save_excel(file_path: Path, content: Dict[str, Any], template_filename: str):
    """
    Excel dosyasını kaydet.
    Şablon dosyayı birebir kopyalar, sadece hücre değerlerini günceller.
    Format/stil bilgileri şablondan korunur.
    """
    template_path = TEMPLATE_DIR / template_filename

    # Şablon dosyayı hedefe kopyala (formatı korumak için)
    if template_path.exists():
        shutil.copy(template_path, file_path)
    else:
        # Şablon yoksa yeni dosya oluştur
        if not file_path.parent.exists():
            file_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        wb.save(file_path)

    # Hedef dosyayı aç ve sadece değerleri yaz
    wb = openpyxl.load_workbook(file_path)
    sheets = content.get("sheets", {})

    for sheet_name, sheet_data in sheets.items():
        if sheet_name not in wb.sheetnames:
            # Şablonda olmayan sheet'i atla
            continue
        ws = wb[sheet_name]

        data = sheet_data.get("data", [])
        
        # Birleştirilmiş hücreleri kontrol et
        merged_ranges = list(ws.merged_cells.ranges)
        
        def is_merged_cell_start(row, col):
            """Birleştirilmiş hücrenin başlangıç hücresi mi kontrol et"""
            for merged_range in merged_ranges:
                if row == merged_range.min_row and col == merged_range.min_col:
                    return True
            return False
        
        def is_in_merged_range(row, col):
            """Hücre birleştirilmiş bir aralığın içinde mi kontrol et"""
            for merged_range in merged_ranges:
                if (merged_range.min_row <= row <= merged_range.max_row and
                    merged_range.min_col <= col <= merged_range.max_col):
                    # Eğer başlangıç hücresi değilse, birleştirilmiş hücrenin içinde
                    if not (row == merged_range.min_row and col == merged_range.min_col):
                        return True
            return False
        
        for row_idx, row_data in enumerate(data, 1):
            for col_idx, cell_data in enumerate(row_data, 1):
                # Birleştirilmiş hücrenin içindeyse ve başlangıç hücresi değilse atla
                if is_in_merged_range(row_idx, col_idx):
                    continue
                
                try:
                    cell_value = cell_data.get("value", "")
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = cell_value
                except Exception as e:
                    # Birleştirilmiş hücre hatası veya başka bir hata varsa atla
                    continue

    # Aktif sheet'i koru
    active_sheet = content.get("active_sheet")
    if active_sheet and active_sheet in wb.sheetnames:
        wb.active = wb[active_sheet]

    wb.save(file_path)

def save_word(file_path: Path, content: Dict[str, Any], template_filename: str = None):
    """
    Word dosyasını kaydet - orijinal formatı koruyarak
    Şablon dosyayı kopyalar ve sadece metin içeriğini günceller
    """
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    # Şablon dosya yolunu belirle
    if template_filename:
        template_path = TEMPLATE_DIR / template_filename
    else:
        template_path = None
    
    # Şablon varsa kopyala, yoksa yeni belge oluştur
    if template_path and template_path.exists():
        # Şablon dosyayı hedefe kopyala (formatı korumak için)
        shutil.copy(template_path, file_path)
        
        # Kopyalanan dosyayı aç
        doc = Document(file_path)
        
        # ÖNCELİKLE: Placeholder varsa onları değiştir (yeni yöntem - format koruma garantili)
        placeholders = content.get("placeholders", {})
        if placeholders:
            replace_placeholders_in_document(doc, placeholders)
            doc.save(file_path)
            return  # Placeholder kullandık, diğer işlemlere gerek yok
        
        # Eğer paragraf veya tablo verisi yoksa, sadece template kopyasını kullan
        paragraphs_data = content.get("paragraphs", [])
        tables_data = content.get("tables", [])
        
        if not paragraphs_data and not tables_data:
            # Hiçbir değişiklik yapma, template'i olduğu gibi kullan
            return
        
        # Şimdi sadece metin içeriğini güncelle (eski yöntem - artık önerilmiyor)
        paragraphs_data = content.get("paragraphs", [])
        tables_data = content.get("tables", [])
        
        # Paragrafları güncelle (TÜM paragraflar, boş olanlar dahil)
        doc_paragraphs = doc.paragraphs  # TÜM paragrafları al
        for i, para_data in enumerate(paragraphs_data):
            if i < len(doc_paragraphs):
                # Mevcut paragrafın metnini güncelle, formatları koru
                p = doc_paragraphs[i]
                new_text = para_data.get("text", "")
                
                # Tüm runs'ları temizle ve yeni metni ekle
                # İlk run'ı kullan ve formatını koru
                if p.runs:
                    # İlk run'ın formatını kullan
                    first_run = p.runs[0]
                    # Tüm runs'ları temizle
                    for run in p.runs:
                        run.text = ""
                    # İlk run'a yeni metni yaz
                    first_run.text = new_text
                else:
                    # Run yoksa metni direkt ekle
                    p.text = new_text
        
        # Tabloları güncelle (sadece metin içeriği)
        for table_idx, table_data in enumerate(tables_data):
            if table_idx < len(doc.tables):
                table = doc.tables[table_idx]
                
                for row_idx, row_data in enumerate(table_data):
                    if row_idx < len(table.rows):
                        row = table.rows[row_idx]
                        
                        for col_idx, cell_data in enumerate(row_data):
                            if col_idx < len(row.cells):
                                cell = row.cells[col_idx]
                                
                                # Cell içeriğini güncelle
                                if isinstance(cell_data, dict):
                                    new_text = cell_data.get("text", "")
                                else:
                                    new_text = str(cell_data)
                                
                                # Cell'in ilk paragrafını güncelle
                                if cell.paragraphs:
                                    cell_para = cell.paragraphs[0]
                                    if cell_para.runs:
                                        # İlk run'ın formatını koru
                                        first_run = cell_para.runs[0]
                                        for run in cell_para.runs:
                                            run.text = ""
                                        first_run.text = new_text
                                    else:
                                        cell_para.text = new_text
                                else:
                                    cell.text = new_text
        
        doc.save(file_path)
        
    else:
        # Şablon yoksa yeni belge oluştur (eski yöntem)
        doc = Document()
        
        # Paragrafları ekle
        paragraphs = content.get("paragraphs", [])
        for para_data in paragraphs:
            para = doc.add_paragraph()
            
            # Paragraph stilini ayarla
            style_name = para_data.get("style", "Normal")
            try:
                para.style = style_name
            except:
                para.style = "Normal"
            
            # Metni ekle
            text = para_data.get("text", "")
            
            # Format bilgileri varsa uygula
            if "format" in para_data:
                format_data = para_data["format"]
                run = para.add_run(text)
                
                # Font stilleri
                if format_data.get("bold"):
                    run.bold = True
                if format_data.get("italic"):
                    run.italic = True
                if format_data.get("underline"):
                    run.underline = True
                
                # Font boyutu
                if format_data.get("font_size"):
                    run.font.size = Pt(format_data["font_size"])
                
                # Font adı
                if format_data.get("font_name"):
                    run.font.name = format_data["font_name"]
                
                # Font rengi
                if format_data.get("font_color"):
                    try:
                        # RGB hex string'den RGBColor'a çevir
                        rgb_str = format_data["font_color"]
                        if len(rgb_str) == 8:  # ARGB formatı
                            rgb_str = rgb_str[2:]  # Alpha'yı atla
                        if len(rgb_str) == 6:
                            r = int(rgb_str[0:2], 16)
                            g = int(rgb_str[2:4], 16)
                            b = int(rgb_str[4:6], 16)
                            run.font.color.rgb = RGBColor(r, g, b)
                    except:
                        pass
            else:
                # Format bilgisi yoksa sadece metni ekle
                para.add_run(text)
            
            # Hizalama
            if "alignment" in para_data:
                align_str = para_data["alignment"]
                if "LEFT" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif "CENTER" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif "RIGHT" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif "JUSTIFY" in align_str:
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Tabloları ekle
        tables = content.get("tables", [])
        for table_data in tables:
            if not table_data or not table_data[0]:
                continue
            
            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            for i, row_data in enumerate(table_data):
                for j, cell_data in enumerate(row_data):
                    # Cell_data bir dict veya string olabilir
                    if isinstance(cell_data, dict):
                        cell_text = cell_data.get("text", "")
                        cell = table.cell(i, j)
                        cell.text = cell_text
                        
                        # Format bilgileri varsa uygula
                        if "format" in cell_data:
                            format_data = cell_data["format"]
                            if cell.paragraphs:
                                para = cell.paragraphs[0]
                                run = para.runs[0] if para.runs else para.add_run(cell_text)
                                
                                # Font stilleri
                                if format_data.get("bold"):
                                    run.bold = True
                                if format_data.get("italic"):
                                    run.italic = True
                                
                                # Font boyutu
                                if format_data.get("font_size"):
                                    run.font.size = Pt(format_data["font_size"])
                                
                                # Font adı
                                if format_data.get("font_name"):
                                    run.font.name = format_data["font_name"]
                                
                                # Font rengi
                                if format_data.get("font_color"):
                                    try:
                                        rgb_str = format_data["font_color"]
                                        if len(rgb_str) == 8:
                                            rgb_str = rgb_str[2:]
                                        if len(rgb_str) == 6:
                                            r = int(rgb_str[0:2], 16)
                                            g = int(rgb_str[2:4], 16)
                                            b = int(rgb_str[4:6], 16)
                                            run.font.color.rgb = RGBColor(r, g, b)
                                    except:
                                        pass
                    else:
                        # Eski format (sadece string)
                        table.cell(i, j).text = str(cell_data)
        
        doc.save(file_path)
def replace_in_paragraph(paragraph, placeholder, value):
    """
    Paragraftaki placeholder'ı değiştir, formatı koru
    """
    if placeholder in paragraph.text:
        # Inline replacement yaparak formatı koruyoruz
        inline = paragraph.runs
        for run in inline:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, str(value))

def replace_placeholders_in_document(doc, placeholders):
    """
    Belgedeki tüm placeholder'ları değiştir
    placeholders: {"MUSTERI_ADI": "ABC Firma", "TARIH": "25.12.2025", ...}
    """
    # Paragrafları güncelle
    for paragraph in doc.paragraphs:
        for key, value in placeholders.items():
            placeholder_text = f"{{{{{key}}}}}"  # {{KEY}} formatı
            if placeholder_text in paragraph.text:
                replace_in_paragraph(paragraph, placeholder_text, value)
    
    # Tabloları güncelle
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for key, value in placeholders.items():
                        placeholder_text = f"{{{{{key}}}}}"
                        if placeholder_text in paragraph.text:
                            replace_in_paragraph(paragraph, placeholder_text, value)
    
    # Header ve Footer'ları güncelle
    for section in doc.sections:
        # Header
        if section.header:
            for paragraph in section.header.paragraphs:
                for key, value in placeholders.items():
                    placeholder_text = f"{{{{{key}}}}}"
                    if placeholder_text in paragraph.text:
                        replace_in_paragraph(paragraph, placeholder_text, value)
        
        # Footer
        if section.footer:
            for paragraph in section.footer.paragraphs:
                for key, value in placeholders.items():
                    placeholder_text = f"{{{{{key}}}}}"
                    if placeholder_text in paragraph.text:
                        replace_in_paragraph(paragraph, placeholder_text, value)
